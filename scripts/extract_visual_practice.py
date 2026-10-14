"""Copy embedded PPT pictures byte-for-byte into their portfolio categories."""

from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r"C:\Users\Lenovo\Desktop\Untitled.pptx")
DEST = ROOT / "public" / "media" / "visual-practice"
CATEGORY = {
    1: "spatial", 2: "spatial", 3: "spatial",
    4: "painting", 5: "photography", 6: "photography",
    7: "graphic-design", 8: "graphic-design", 9: "graphic-design", 10: "graphic-design",
}


def pictures(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from pictures(shape.shapes)
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            yield shape


presentation = Presentation(SOURCE)
manifest = {"source": SOURCE.name, "slides": []}
for slide_number, slide in enumerate(presentation.slides, 1):
    folder = DEST / CATEGORY[slide_number]
    folder.mkdir(parents=True, exist_ok=True)
    items = []
    for picture_number, shape in enumerate(pictures(slide.shapes), 1):
        image = shape.image
        data = image.blob
        dimensions = Image.open(BytesIO(data)).size
        name = f"slide-{slide_number:02d}-picture-{picture_number:03d}.{image.ext}"
        target = folder / name
        target.write_bytes(data)
        items.append({
            "file": f"/media/visual-practice/{CATEGORY[slide_number]}/{name}",
            "width": dimensions[0], "height": dimensions[1],
            "sha256": sha256(data).hexdigest(),
        })
    manifest["slides"].append({"number": slide_number, "category": CATEGORY[slide_number], "pictures": items})

(DEST / "source-manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
print("slides:", len(manifest["slides"]), "pictures:", sum(len(s["pictures"]) for s in manifest["slides"]))
